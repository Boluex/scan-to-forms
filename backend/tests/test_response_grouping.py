import csv
from io import BytesIO, StringIO
from unittest.mock import patch

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfWriter

from apps.billing.models import Plan, UsageRecord
from apps.documents.models import DocumentPage, OCRJob, OCRRegion, OCRResult, UploadedDocument
from apps.documents.services.grouping import (
    aggregate_response_answers,
    classify_page,
    create_document_pages,
    create_response,
    refresh_response_validation,
)
from apps.documents.services.ocr import DocumentOutput, PageOutput
from apps.documents.tasks import process_document
from apps.documents.views import UploadThrottle
from apps.googleforms.models import AppsScriptJob
from apps.googleforms.services import resolve_script_source
from apps.questionnaires.models import (
    Answer,
    Question,
    Questionnaire,
    QuestionnaireVersion,
    Response,
    ResponseBatch,
    TemplatePage,
)


@pytest.fixture
def generous_free_plan(db, monkeypatch):
    monkeypatch.setattr(UploadThrottle, "rate", "10000/hour", raising=False)
    plan, _ = Plan.objects.get_or_create(
        code=Plan.Code.FREE,
        defaults={
            "name": "Free",
            "monthly_price_kobo": 0,
            "monthly_page_limit": 1000,
            "batch_size_limit": 500,
            "max_team_members": 1,
        },
    )
    plan.monthly_page_limit = 1000
    plan.batch_size_limit = 500
    plan.has_xlsx = True
    plan.save(update_fields=("monthly_page_limit", "batch_size_limit", "has_xlsx", "updated_at"))
    return plan


def study(user, expected_pages=4):
    questionnaire = Questionnaire.objects.create(owner=user, title="Student Housing Survey")
    version = QuestionnaireVersion.objects.create(
        questionnaire=questionnaire,
        version_number=1,
        expected_page_count=expected_pages,
    )
    batch = ResponseBatch.objects.create(owner=user, questionnaire_version=version, name="August Physical Responses")
    return questionnaire, version, batch


def pdf_upload(name, page_count):
    output = BytesIO()
    writer = PdfWriter()
    for _ in range(page_count):
        writer.add_blank_page(width=595, height=842)
    writer.write(output)
    return SimpleUploadedFile(name, output.getvalue(), content_type="application/pdf")


def image_upload(name, marker=0):
    output = BytesIO()
    Image.new("RGB", (32, 32), (marker % 251, (marker * 7) % 251, (marker * 13) % 251)).save(output, "PNG")
    return SimpleUploadedFile(name, output.getvalue(), content_type="image/png")


def post_document(client, batch, upload, **extra):
    payload = {"document_type": "RESPONSE", "batch": str(batch.id), "upload": upload, **extra}
    return client.post("/api/v1/documents/", payload, format="multipart")


@pytest.mark.django_db(transaction=True)
def test_one_four_page_pdf_becomes_one_response(client, user, settings, tmp_path, generous_free_plan):
    settings.MEDIA_ROOT = tmp_path
    _, _, batch = study(user)
    with patch("apps.documents.views.process_document.delay"):
        uploaded = post_document(client, batch, pdf_upload("respondent_001.pdf", 4))
    assert uploaded.status_code == 201
    response = Response.objects.get(batch=batch)
    assert response.pages.count() == 4
    assert response.status == Response.Status.READY_FOR_PROCESSING
    assert list(response.pages.values_list("assigned_template_page_number", flat=True)) == [1, 2, 3, 4]


@pytest.mark.django_db(transaction=True)
def test_two_four_page_pdfs_become_two_responses(client, user, settings, tmp_path, generous_free_plan):
    settings.MEDIA_ROOT = tmp_path
    _, _, batch = study(user)
    with patch("apps.documents.views.process_document.delay"):
        for number in (1, 2):
            uploaded = post_document(client, batch, pdf_upload(f"respondent_{number:03d}.pdf", 4))
            assert uploaded.status_code == 201
    assert batch.responses.count() == 2
    assert list(batch.responses.values_list("pages__id", flat=True)).count(None) == 0
    assert [response.pages.count() for response in batch.responses.all()] == [4, 4]


def upload_prepared_images(client, batch, assignments, start_marker=1):
    for assignment in assignments:
        index = assignment["upload_index"]
        uploaded = post_document(
            client,
            batch,
            image_upload(f"IMG_{index:04d}.png", start_marker + index),
            response=assignment["response_id"],
            template_page_number=str(assignment["template_page_number"]),
            grouping_mode="BULK_ORDERED",
        )
        assert uploaded.status_code == 201, uploaded.data


@pytest.mark.django_db(transaction=True)
def test_eight_ordered_images_become_two_responses(client, user, settings, tmp_path, generous_free_plan):
    settings.MEDIA_ROOT = tmp_path
    _, _, batch = study(user)
    prepared = client.post(f"/api/v1/response-batches/{batch.id}/prepare-bulk/", {"file_count": 8}, format="json")
    assert prepared.status_code == 201
    with patch("apps.documents.views.process_document.delay"):
        upload_prepared_images(client, batch, prepared.data["assignments"])
    assert batch.responses.count() == 2
    assert [response.pages.count() for response in batch.responses.all()] == [4, 4]


@pytest.mark.django_db(transaction=True)
def test_two_hundred_ordered_images_become_fifty_responses(client, user, settings, tmp_path, generous_free_plan):
    settings.MEDIA_ROOT = tmp_path
    _, _, batch = study(user)
    prepared = client.post(f"/api/v1/response-batches/{batch.id}/prepare-bulk/", {"file_count": 200}, format="json")
    assert prepared.status_code == 201
    assert prepared.data["response_count"] == 50
    with patch("apps.documents.views.process_document.delay"):
        upload_prepared_images(client, batch, prepared.data["assignments"], start_marker=20)
    assert batch.responses.count() == 50
    assert DocumentPage.objects.filter(response__batch=batch).count() == 200


def manual_response(client, batch, reference=""):
    created = client.post(
        "/api/v1/responses/",
        {"batch": str(batch.id), "respondent_reference": reference},
        format="json",
    )
    assert created.status_code == 201, created.data
    return Response.objects.get(pk=created.data["id"])


def upload_manual_pages(client, batch, response, assignments):
    for upload_order, page_number in enumerate(assignments, start=1):
        uploaded = post_document(
            client,
            batch,
            image_upload(f"manual-{upload_order}-{page_number}.png", upload_order * 19 + page_number),
            response=str(response.id),
            template_page_number=str(page_number),
            grouping_mode="MANUAL_RESPONSE",
        )
        assert uploaded.status_code == 201, uploaded.data


@pytest.mark.django_db(transaction=True)
def test_missing_page_is_detected(client, user, settings, tmp_path, generous_free_plan):
    settings.MEDIA_ROOT = tmp_path
    _, _, batch = study(user)
    response = manual_response(client, batch)
    with patch("apps.documents.views.process_document.delay"):
        upload_manual_pages(client, batch, response, [1, 2, 4])
    response.refresh_from_db()
    assert response.status == Response.Status.INCOMPLETE
    assert response.validation_issues[0]["code"] == "MISSING_PAGES"
    assert response.validation_issues[0]["pages"] == [3]


@pytest.mark.django_db(transaction=True)
def test_duplicate_page_and_missing_page_are_detected(client, user, settings, tmp_path, generous_free_plan):
    settings.MEDIA_ROOT = tmp_path
    _, _, batch = study(user)
    response = manual_response(client, batch)
    with patch("apps.documents.views.process_document.delay"):
        upload_manual_pages(client, batch, response, [1, 2, 2, 4])
    response.refresh_from_db()
    codes = {issue["code"] for issue in response.validation_issues}
    assert {"MISSING_PAGES", "DUPLICATE_PAGES"} <= codes
    assert response.status == Response.Status.NEEDS_REVIEW
    assert response.pages.filter(assigned_template_page_number=2).count() == 2


@pytest.mark.django_db(transaction=True)
def test_out_of_order_pages_are_logically_reordered(client, user, settings, tmp_path, generous_free_plan):
    settings.MEDIA_ROOT = tmp_path
    _, _, batch = study(user)
    response = manual_response(client, batch)
    with patch("apps.documents.views.process_document.delay"):
        upload_manual_pages(client, batch, response, [1, 3, 2, 4])
    detail = client.get(f"/api/v1/responses/{response.id}/")
    assert detail.status_code == 200
    assert [page["assigned_template_page_number"] for page in detail.data["pages"]] == [1, 2, 3, 4]
    assert [page["original_upload_order"] for page in detail.data["pages"]] == [1, 3, 2, 4]
    assert any(issue["code"] == "OUT_OF_ORDER" for issue in detail.data["validation_issues"])


@pytest.mark.django_db
def test_page_marker_classifier_corrects_provisional_order(user):
    _, _, batch = study(user)
    response = create_response(batch)
    document = UploadedDocument.objects.create(
        owner=user,
        batch=batch,
        response=response,
        document_type=UploadedDocument.Type.RESPONSE,
        page_count=4,
        original_filename="out-of-order.pdf",
        content_type="application/pdf",
        size_bytes=1,
        sha256="f" * 64,
    )
    pages = create_document_pages(document)
    for page, marker in zip(pages, (1, 3, 2, 4), strict=True):
        classify_page(page, f"Student Housing Survey — Page {marker} of 4")
    refresh_response_validation(response.id)
    assert list(response.pages.order_by("original_upload_order").values_list("assigned_template_page_number", flat=True)) == [1, 3, 2, 4]
    assert list(response.pages.values_list("assigned_template_page_number", flat=True)) == [1, 2, 3, 4]


@pytest.mark.django_db
def test_template_text_similarity_classifies_a_page_without_a_page_marker(user):
    _, version, batch = study(user, expected_pages=2)
    TemplatePage.objects.create(
        version=version,
        page_number=1,
        reference_text="Student Housing Survey personal profile age gender department",
    )
    TemplatePage.objects.create(
        version=version,
        page_number=2,
        reference_text="Student Housing Survey rent utilities landlord satisfaction recommendations",
    )
    response = create_response(batch)
    document = UploadedDocument.objects.create(
        owner=user,
        batch=batch,
        response=response,
        document_type=UploadedDocument.Type.RESPONSE,
        page_count=1,
        original_filename="unknown-page.png",
        content_type="image/png",
        size_bytes=1,
        sha256="a" * 64,
    )
    page = create_document_pages(document)[0]
    classify_page(page, "Rent and utilities paid to landlord. Satisfaction and recommendations.")
    page.refresh_from_db()
    assert page.detected_template_page_number == 2
    assert page.assigned_template_page_number == 2
    assert page.classification_method == DocumentPage.ClassificationMethod.TEXT_SIMILARITY
    assert page.classification_confidence >= .35


def completed_page(user, response, page_number, question_text, answer_text):
    document = UploadedDocument.objects.create(
        owner=user,
        batch=response.batch,
        response=response,
        document_type=UploadedDocument.Type.RESPONSE,
        grouping_mode=UploadedDocument.GroupingMode.MANUAL_RESPONSE,
        page_count=1,
        original_filename=f"page-{page_number}.png",
        content_type="image/png",
        size_bytes=1,
        sha256=f"{page_number:064x}",
    )
    page = DocumentPage.objects.create(
        document=document,
        response=response,
        page_number=1,
        original_upload_order=page_number,
        assigned_template_page_number=page_number,
        classification_method=DocumentPage.ClassificationMethod.MANUAL,
        classification_confidence=1,
        processing_status=DocumentPage.ProcessingStatus.COMPLETED,
    )
    job = OCRJob.objects.create(document=document, status=OCRJob.Status.NEEDS_REVIEW)
    result = OCRResult.objects.create(job=job, page=page, plain_text=f"{question_text}\n{answer_text}", mean_confidence=.95)
    OCRRegion.objects.create(result=result, sequence=0, text=question_text, confidence=.98, bounding_box={})
    OCRRegion.objects.create(result=result, sequence=1, text=answer_text, confidence=.95, bounding_box={})
    return page


@pytest.mark.django_db
def test_answers_from_multiple_pages_merge_into_one_response(user):
    _, version, batch = study(user)
    questions = []
    for page_number in range(1, 5):
        questions.append(
            Question.objects.create(
                version=version,
                key=f"q{page_number}",
                position=page_number,
                text=f"Question {page_number}?",
                type=Question.Type.SHORT_TEXT,
                template_page_number=page_number,
            )
        )
    response = create_response(batch)
    for page_number, question in enumerate(questions, start=1):
        completed_page(user, response, page_number, question.text, f"Answer {page_number}")
    response, aggregated = aggregate_response_answers(response.id, force=True)
    assert aggregated is True
    assert response.answers.count() == 4
    assert list(response.answers.values_list("value_text", flat=True)) == ["Answer 1", "Answer 2", "Answer 3", "Answer 4"]
    assert {answer.source_region["template_page_number"] for answer in response.answers.all()} == {1, 2, 3, 4}


@pytest.mark.django_db
def test_csv_and_xlsx_have_one_row_per_multi_page_respondent(client, user, generous_free_plan):
    _, version, batch = study(user)
    question = Question.objects.create(version=version, key="q1", position=1, text="Department", type="SHORT_TEXT")
    for sequence in range(1, 3):
        response = create_response(batch, f"R-{sequence:03d}")
        for page_number in range(1, 5):
            completed_page(user, response, page_number, "Unused?", "Unused")
        response.status = Response.Status.CONFIRMED
        response.save(update_fields=("status", "updated_at"))
        Answer.objects.create(response=response, question=question, value_text=f"Department {sequence}", confidence=.99)

    csv_response = client.get(f"/api/v1/exports/batches/{batch.id}/csv/")
    rows = list(csv.reader(StringIO(csv_response.content.decode("utf-8-sig"))))
    assert len(rows) == 3
    xlsx_response = client.get(f"/api/v1/exports/batches/{batch.id}/xlsx/")
    workbook = load_workbook(BytesIO(xlsx_response.content), read_only=True)
    assert sum(1 for _ in workbook["Responses"].iter_rows(values_only=True)) == 3


@pytest.mark.django_db(transaction=True)
def test_usage_counts_physical_pages_not_respondents(client, user, settings, tmp_path, generous_free_plan):
    settings.MEDIA_ROOT = tmp_path
    _, _, batch = study(user)
    with patch("apps.documents.views.process_document.delay"):
        uploaded = post_document(client, batch, pdf_upload("four-pages.pdf", 4))
    assert uploaded.status_code == 201
    usage = UsageRecord.objects.get(account=user)
    assert usage.pages_used == 4
    assert batch.responses.count() == 1


@pytest.mark.django_db
def test_other_user_cannot_access_response_page(client, other_user):
    _, _, batch = study(other_user)
    response = create_response(batch)
    document = UploadedDocument.objects.create(
        owner=other_user,
        batch=batch,
        response=response,
        document_type=UploadedDocument.Type.RESPONSE,
        original_filename="private.png",
        content_type="image/png",
        size_bytes=1,
        sha256="e" * 64,
    )
    page = create_document_pages(document)[0]
    assert client.get(f"/api/v1/response-pages/{page.id}/").status_code == 404
    assert client.patch(f"/api/v1/response-pages/{page.id}/", {"assigned_template_page_number": 1}).status_code == 404


@pytest.mark.django_db
def test_manual_image_page_can_move_to_another_response_and_old_answers_are_cleared(client, user):
    _, version, batch = study(user, expected_pages=1)
    question = Question.objects.create(
        version=version,
        key="q1",
        position=1,
        text="Department?",
        type=Question.Type.SHORT_TEXT,
        template_page_number=1,
    )
    source = create_response(batch)
    target = create_response(batch)
    page = completed_page(user, source, 1, question.text, "Computer Science")
    source, aggregated = aggregate_response_answers(source.id, force=True)
    assert aggregated is True
    assert source.answers.count() == 1

    moved = client.patch(
        f"/api/v1/response-pages/{page.id}/",
        {"response": str(target.id), "assigned_template_page_number": 1},
        format="json",
    )
    assert moved.status_code == 200, moved.data
    page.refresh_from_db()
    page.document.refresh_from_db()
    source.refresh_from_db()
    target.refresh_from_db()
    assert page.response_id == target.id
    assert page.document.response_id == target.id
    assert source.answers.count() == 0
    assert source.status == Response.Status.INCOMPLETE
    assert target.answers.count() == 1


@pytest.mark.django_db
def test_manual_page_number_override_is_authoritative(client, user):
    _, _, batch = study(user, expected_pages=2)
    response = create_response(batch)
    first = completed_page(user, response, 1, "Question one?", "Answer one")
    completed_page(user, response, 2, "Question two?", "Answer two")
    first.detected_template_page_number = 2
    first.assigned_template_page_number = 2
    first.classification_method = DocumentPage.ClassificationMethod.PAGE_MARKER
    first.classification_confidence = .99
    first.save(
        update_fields=(
            "detected_template_page_number",
            "assigned_template_page_number",
            "classification_method",
            "classification_confidence",
            "updated_at",
        )
    )
    refresh_response_validation(response.id)

    corrected = client.patch(
        f"/api/v1/response-pages/{first.id}/",
        {"response": str(response.id), "assigned_template_page_number": 1},
        format="json",
    )
    assert corrected.status_code == 200, corrected.data
    first.refresh_from_db()
    response.refresh_from_db()
    assert first.classification_method == DocumentPage.ClassificationMethod.MANUAL
    assert first.assigned_template_page_number == 1
    assert not any(issue["code"] == "PAGE_MISMATCH" for issue in first.validation_issues)
    assert not any(issue["code"] in {"MISSING_PAGES", "DUPLICATE_PAGES"} for issue in response.validation_issues)


@pytest.mark.django_db
def test_google_forms_source_has_one_record_for_one_four_page_respondent(user):
    _, version, batch = study(user)
    question = Question.objects.create(
        version=version,
        key="department",
        position=1,
        text="Department?",
        type=Question.Type.SHORT_TEXT,
        template_page_number=1,
    )
    response = create_response(batch)
    for page_number in range(1, 5):
        completed_page(user, response, page_number, "Unused?", "Unused")
    Answer.objects.create(response=response, question=question, value_text="Computer Science", confidence=.99)
    response.status = Response.Status.CONFIRMED
    response.save(update_fields=("status", "updated_at"))

    source = resolve_script_source(user, AppsScriptJob.SourceType.RESPONSE_BATCH, batch.id)
    assert len(source.responses) == 1
    assert source.responses[0]["id"] == str(response.id)
    assert source.responses[0]["answers"] == {"department": "Computer Science"}


@pytest.mark.django_db
def test_failed_ocr_page_prevents_response_success(user):
    _, version, batch = study(user)
    Question.objects.create(version=version, key="q1", position=1, text="Question?", type="SHORT_TEXT")
    response = create_response(batch)
    for page_number in range(1, 5):
        page = completed_page(user, response, page_number, "Question?", "Answer")
        if page_number == 3:
            page.processing_status = DocumentPage.ProcessingStatus.FAILED
            page.save(update_fields=("processing_status", "updated_at"))
            page.document.ocr_job.status = OCRJob.Status.FAILED
            page.document.ocr_job.save(update_fields=("status", "updated_at"))
    response, aggregated = aggregate_response_answers(response.id, force=True)
    assert aggregated is False
    assert response.status == Response.Status.FAILED
    assert response.answers.count() == 0


@pytest.mark.django_db(transaction=True)
def test_ocr_output_missing_one_pdf_page_fails_the_response(
    client,
    user,
    settings,
    tmp_path,
    generous_free_plan,
):
    settings.MEDIA_ROOT = tmp_path
    _, _, batch = study(user)
    with patch("apps.documents.views.process_document.delay"):
        uploaded = post_document(client, batch, pdf_upload("incomplete-ocr.pdf", 4))
    assert uploaded.status_code == 201
    document = UploadedDocument.objects.get(pk=uploaded.data["id"])
    partial_output = DocumentOutput(
        engine="test-engine",
        model_version="test-1",
        pages=[
            PageOutput(
                page_number=number,
                width=595,
                height=842,
                text=f"Page {number} of 4",
                confidence=.9,
                regions=[],
                preprocessing={},
            )
            for number in range(1, 4)
        ],
    )
    with patch("apps.documents.tasks.extract_document", return_value=partial_output):
        with pytest.raises(RuntimeError, match="source page.*4"):
            process_document.run(str(document.ocr_job.id))

    document.refresh_from_db()
    document.response.refresh_from_db()
    assert document.ocr_job.status == OCRJob.Status.FAILED
    assert document.response.status == Response.Status.FAILED
    assert document.response.pages.filter(processing_status=DocumentPage.ProcessingStatus.FAILED).count() == 4
    assert document.response.answers.count() == 0


@pytest.mark.django_db(transaction=True)
def test_existing_single_page_upload_still_creates_one_response(client, user, settings, tmp_path, generous_free_plan):
    settings.MEDIA_ROOT = tmp_path
    _, _, batch = study(user, expected_pages=1)
    with patch("apps.documents.views.process_document.delay"):
        uploaded = post_document(client, batch, image_upload("single.png", 77))
    assert uploaded.status_code == 201
    response = batch.responses.get()
    assert response.pages.count() == 1
    assert response.status == Response.Status.READY_FOR_PROCESSING
